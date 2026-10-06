#!/usr/bin/env python3
"""Diapositivas para Diagnóstico y Tuning de BD asistido por IA."""

from pptx.util import Inches, Pt

from estilo_base import (
    ACCENT_BLUE,
    ACCENT_SUBTLE,
    TEXT_TITLE,
    FONT_NAME,
    add_card,
    add_structured_item,
    create_base_slide,
    create_empty_deck,
    save_deck_safe,
)


def _add_two_cards(prs, section, title, left_title, left_items, right_title, right_items):
    slide = create_base_slide(prs, section, title)
    for x, card_title, items in (
        (0.8, left_title, left_items),
        (6.8, right_title, right_items),
    ):
        add_card(slide, Inches(x), Inches(1.65), Inches(5.7), Inches(5.1), card_title)
        text_box = slide.shapes.add_textbox(
            Inches(x + 0.25), Inches(2.25), Inches(5.2), Inches(4.25)
        )
        text_frame = text_box.text_frame
        text_frame.word_wrap = True
        for label, detail in items:
            add_structured_item(text_frame, label, detail, 12)
    return slide


def _add_three_cards(prs, section, title, cards):
    slide = create_base_slide(prs, section, title)
    positions = (0.8, 4.8, 8.8)
    for x, (card_title, items) in zip(positions, cards):
        add_card(slide, Inches(x), Inches(1.65), Inches(3.7), Inches(5.1), card_title)
        text_box = slide.shapes.add_textbox(
            Inches(x + 0.2), Inches(2.25), Inches(3.3), Inches(4.25)
        )
        text_frame = text_box.text_frame
        text_frame.word_wrap = True
        for label, detail in items:
            add_structured_item(text_frame, label, detail, 11)
    return slide


def agregar_slides_diagnostico_ia(prs):
    """Añade diez diapositivas para la exposición y la demostración."""
    cover = create_base_slide(
        prs,
        "Arquitectura de Datos e Inteligencia Artificial",
        "Diagnóstico y Tuning de Bases de Datos Asistido por IA",
    )
    subtitle = cover.shapes.add_textbox(Inches(0.85), Inches(1.65), Inches(11.5), Inches(0.8))
    subtitle.text_frame.word_wrap = True
    subtitle_paragraph = subtitle.text_frame.paragraphs[0]
    subtitle_paragraph.text = "Integración con OpenMetadata y operación bajo un modelo MSP"
    subtitle_paragraph.font.name = FONT_NAME
    subtitle_paragraph.font.size = Pt(20)
    subtitle_paragraph.font.color.rgb = ACCENT_SUBTLE

    add_card(cover, Inches(0.8), Inches(3.05), Inches(11.7), Inches(2.7), "Pregunta central")
    cover_text = cover.shapes.add_textbox(Inches(1.1), Inches(3.8), Inches(11.0), Inches(1.5))
    cover_text.text_frame.word_wrap = True
    cover_paragraph = cover_text.text_frame.paragraphs[0]
    cover_paragraph.text = (
        "¿Cómo convertir señales de PostgreSQL en recomendaciones explicables, "
        "seguras y operables?"
    )
    cover_paragraph.font.name = FONT_NAME
    cover_paragraph.font.size = Pt(24)
    cover_paragraph.font.bold = True
    cover_paragraph.font.color.rgb = TEXT_TITLE

    _add_two_cards(
        prs,
        "01 · Contexto",
        "Del tuning reactivo al diagnóstico asistido",
        "Enfoque tradicional",
        [
            ("Reacción:", "El análisis suele empezar después de una alerta o queja de latencia."),
            ("Correlación manual:", "Logs, planes y métricas viven en herramientas diferentes."),
            ("Costo:", "El DBA invierte tiempo en localizar y reproducir el problema."),
            ("Riesgo:", "La degradación continúa mientras se recopila evidencia."),
        ],
        "IA como asistente",
        [
            ("Prioriza:", "Agrupa consultas y ordena anomalías por impacto probable."),
            ("Correlaciona:", "Relaciona SQL, plan, métricas y cambios en el tiempo."),
            ("Anticipa:", "Usa históricos para detectar tendencias y regresiones."),
            ("No reemplaza:", "El DBA valida causa, riesgo y acción antes del cambio."),
        ],
    )

    _add_two_cards(
        prs,
        "02 · Evidencia técnica",
        "¿Qué señales puede analizar sin conocer la BD?",
        "Telemetría de entrada",
        [
            ("Slow Query Logs:", "Frecuencia, duración y patrones de SQL normalizado."),
            ("Execution Plans:", "EXPLAIN ANALYZE muestra nodos, filas y tiempos observados."),
            ("Buffers:", "Distingue páginas en caché de bloques leídos durante el plan."),
            ("Wait events:", "Ayudan a separar espera por I/O, locks, CPU u otros recursos."),
        ],
        "Inferencias y límites",
        [
            ("Profiling:", "pg_stat_statements prioriza consultas por costo y frecuencia."),
            ("Índices:", "Un filtro selectivo sin índice puede sugerir un candidato."),
            ("Contexto necesario:", "Esquema, cardinalidad y estadísticas mejoran la hipótesis."),
            ("Límite:", "La IA no infiere semántica ni distribución que no se le entregue."),
        ],
    )

    _add_two_cards(
        prs,
        "03 · Catálogo y gobierno",
        "OpenMetadata: contexto, no motor de tuning",
        "Lo que SÍ aporta",
        [
            ("Catálogo:", "Inventario de bases, esquemas, tablas y columnas."),
            ("Linaje:", "Relaciones entre fuentes, tablas, pipelines y consumidores."),
            ("Gobierno:", "Propiedad, clasificación, descripciones y trazabilidad."),
            ("Contexto para IA:", "El agente consulta columnas, descripción y tags por API."),
        ],
        "Lo que NO hace",
        [
            ("No ejecuta consultas:", "PostgreSQL sigue planificando y ejecutando SQL."),
            ("No es observabilidad:", "No sustituye logs, EXPLAIN ni métricas de runtime."),
            ("No crea índices:", "No decide si el costo de escritura es aceptable."),
            ("Estado del laboratorio:", "OpenMetadata 2.0.3 + ingesta PostgreSQL en Compose separado."),
        ],
    )

    _add_two_cards(
        prs,
        "04 · Arquitectura",
        "Stack local con catálogo y diagnóstico integrado",
        "Servicios desplegables",
        [
            ("PostgreSQL 16:", "Esquema sintético y extensión pg_stat_statements."),
            ("Promtail:", "Lee los logs del contenedor en modo solo lectura."),
            ("Loki + Grafana:", "Centralizan y muestran eventos consultables."),
            ("OpenMetadata:", "Ingiere y cataloga el esquema PostgreSQL."),
        ],
        "Diagnóstico y control",
        [
            ("1. Evidencia:", "El agente obtiene plan real con EXPLAIN JSON."),
            ("2. Contexto:", "Consulta FQN, columnas, descripción y tags por API."),
            ("3. Análisis:", "Ollama/OpenAI devuelve una hipótesis contrastable."),
            ("4. Gobierno:", "DBA/MSP aprueba, prueba y mide el índice."),
        ],
    )

    _add_three_cards(
        prs,
        "05 · Servicio gestionado",
        "Operación de base de datos bajo un modelo MSP",
        [
            (
                "Operación 24/7",
                [
                    ("Cobertura:", "Monitoreo, guardias y escalamiento por severidad."),
                    ("Continuidad:", "Backups, restauración, capacidad y parchado."),
                    ("Tuning:", "Revisión de latencia y cambios de forma continua."),
                ],
            ),
            (
                "Acuerdos medibles",
                [
                    ("SLA:", "Compromiso contractual acordado con el cliente."),
                    ("SLO:", "Objetivo operativo medible del servicio."),
                    ("SLI:", "Indicador observado, por ejemplo disponibilidad."),
                ],
            ),
            (
                "Control de cambios",
                [
                    ("IA:", "Clasifica y prepara diagnóstico con evidencia."),
                    ("MSP/DBA:", "Valida impacto, ventana, aprobación y rollback."),
                    ("Cliente:", "Define criticidad, políticas y aceptación de riesgo."),
                ],
            ),
        ],
    )

    _add_two_cards(
        prs,
        "06 · Laboratorio",
        "Caso reproducible sobre PostgreSQL del proyecto",
        "Línea base",
        [
            ("Entorno:", "Docker Compose; PostgreSQL 16 y datos sintéticos existentes."),
            ("Tabla:", "transacciones, con cientos de miles de filas en volumen inicial."),
            ("Consulta:", "Historial de una cuenta con filtro temporal, orden y límite."),
            ("Medición:", "EXPLAIN (ANALYZE, BUFFERS, FORMAT JSON)."),
        ],
        "Consulta del experimento",
        [
            ("Filtro:", "cuenta_id = 1520 y fecha_hora desde 2024-01-01."),
            ("Problema buscado:", "Recorrido amplio y trabajo de ordenamiento."),
            ("Salida al agente:", "Plan JSON y columnas leídas desde information_schema."),
            ("Cautela:", "ANALYZE ejecuta la consulta; no usar a ciegas en producción."),
        ],
    )

    _add_two_cards(
        prs,
        "07 · Recomendación y verificación",
        "Del diagnóstico a una mejora medida",
        "Agente IA",
        [
            ("Entrada:", "SQL, plan observado y columnas reales de la tabla."),
            ("Diagnóstico:", "Explica el cuello con evidencia, no solo con intuición."),
            ("DDL candidato:", "CREATE INDEX ON transacciones (cuenta_id, fecha_hora DESC)."),
            ("Reescritura:", "Propone alternativas solo si preservan el resultado."),
        ],
        "Validación humana",
        [
            ("Comparar:", "Nodos del plan, filas, buffers y mediana de ejecución."),
            ("Revisar costo:", "Espacio, escritura adicional, mantenimiento y bloqueos."),
            ("Aplicar:", "El lab ejecuta un DDL fijo; nunca SQL libre de la IA."),
            ("Operar:", "Desplegar con aprobación, monitorizar y tener rollback."),
        ],
    )

    _add_three_cards(
        prs,
        "08 · Demo en vivo",
        "Secuencia de exposición: evidencia, recomendación, control",
        [
            (
                "Medir",
                [
                    ("Comando:", "docker compose up -d"),
                    ("Después:", "py .\\scripts\\ai_db_tuning.py --use-openmetadata"),
                    ("Mostrar:", "Plan inicial, scan, buffers y estimaciones."),
                ],
            ),
            (
                "Analizar",
                [
                    ("Modelo:", "Ollama local o OpenAI API configurado."),
                    ("Explicar:", "La IA recibe evidencia y esquema, no magia."),
                    ("Gobernar:", "Se inspecciona recomendación antes de cambiar."),
                ],
            ),
            (
                "Verificar",
                [
                    ("Comando:", "py .\\scripts\\ai_db_tuning.py --use-openmetadata --apply-demo-index"),
                    ("Contrastar:", "Plan posterior, buffers y tiempo mediano."),
                    ("Cerrar:", "MSP formaliza aprobación, SLO y aprendizaje."),
                ],
            ),
        ],
    )

    closing = create_base_slide(prs, "Conclusión", "La IA acelera el diagnóstico; la responsabilidad sigue siendo humana")
    add_card(closing, Inches(0.8), Inches(1.65), Inches(11.7), Inches(5.1), "Tres ideas para llevar")
    text_box = closing.shapes.add_textbox(Inches(1.1), Inches(2.3), Inches(11.1), Inches(4.1))
    text_frame = text_box.text_frame
    text_frame.word_wrap = True
    add_structured_item(
        text_frame,
        "Evidencia primero:",
        "Logs, planes, métricas y profiling sustentan el diagnóstico.",
        16,
    )
    add_structured_item(
        text_frame,
        "Contexto después:",
        "OpenMetadata añade semántica y linaje, no tuning de bajo nivel.",
        16,
    )
    add_structured_item(
        text_frame,
        "Operación gobernada:",
        "Un MSP mide SLAs/SLOs y el DBA aprueba cambios, mide impacto y mantiene rollback.",
        16,
    )


if __name__ == "__main__":
    preview = create_empty_deck()
    agregar_slides_diagnostico_ia(preview)
    output = "Presentacion_Diagnostico_IA_OpenMetadata_MSP.pptx"
    save_deck_safe(preview, output)