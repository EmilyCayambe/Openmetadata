#!/usr/bin/env python3
"""
slides_optimizacion.py
=============================================================================
MÓDULO 3: OPTIMIZACIÓN DE CONSULTAS SQL Y BENCHMARKS (EXPLAIN ANALYZE)
Responsable: Estudiante 3
=============================================================================
Contenido:
- Slide 8: Afinamiento de Rendimiento: Análisis de Planes y Métricas de E/S
- Slide 9: Resultados Comparativos: De Sequential Scan a Index Scan
- Generación de gráfico analítico con benchmarks medidos en caliente.

Este archivo se puede ejecutar de forma INDEPENDIENTE para generar una vista
previa exclusiva de este módulo:
    python presentacion/slides_optimizacion.py
"""

import os
from pptx.util import Inches
import matplotlib.pyplot as plt
import numpy as np

from estilo_base import (
    ASSETS_DIR, SCRIPT_DIR, create_empty_deck, create_base_slide,
    add_card, add_structured_item, save_deck_safe
)

def generar_grafico_optimizacion():
    """Genera gráfico de barras: Rendimiento Antes vs Después de Índices"""
    path = os.path.join(ASSETS_DIR, "grafico_optimizacion_benchmark.png")
    fig, ax = plt.subplots(figsize=(6.4, 4.0), facecolor='#1E293B')
    ax.set_facecolor('#1E293B')

    casos = ['Búsqueda Puntual\n(B-Tree DNI)', 'Rango y Sort\n(Compuesto Cuentas)', 'Anti-patrón LOWER\n(Índice Funcional)']
    sin_indice = [13.71, 391.04, 71.17]
    con_indice = [0.12, 0.10, 0.72]

    x = np.arange(len(casos))
    width = 0.32

    rects1 = ax.bar(x - width/2, sin_indice, width, label='Sin Índice (Sequential Scan)',
                    color='#F87171', edgecolor='#1E293B', zorder=3)
    rects2 = ax.bar(x + width/2, con_indice, width, label='Con Índice Optimizado (Index Scan)',
                    color='#38BDF8', edgecolor='#1E293B', zorder=3)

    ax.set_title("Impacto de Optimización: Latencia de Ejecución (ms)", color='#F8FAFC',
                 fontsize=11.5, fontweight='bold', pad=14, fontfamily='sans-serif')
    ax.set_ylabel("Tiempo de Ejecución (ms)", color='#94A3B8', fontsize=9.5)
    ax.set_xticks(x)
    ax.set_xticklabels(casos, color='#E2E8F0', fontsize=8.5)
    ax.tick_params(colors='#E2E8F0', labelsize=8.5)
    ax.grid(True, axis='y', linestyle=':', alpha=0.25, color='#64748B')

    for rect in rects1:
        h = rect.get_height()
        ax.annotate(f'{h:.1f} ms', xy=(rect.get_x() + rect.get_width() / 2, h),
                    xytext=(0, 3), textcoords="offset points", ha='center', va='bottom',
                    color='#FCA5A5', fontsize=8, fontweight='bold')

    for rect in rects2:
        h = rect.get_height()
        ax.annotate(f'{h:.2f} ms', xy=(rect.get_x() + rect.get_width() / 2, h),
                    xytext=(0, 3), textcoords="offset points", ha='center', va='bottom',
                    color='#7DD3FC', fontsize=8, fontweight='bold')

    for spine in ax.spines.values():
        spine.set_color('#334155')

    legend = ax.legend(facecolor='#0F172A', edgecolor='#334155', fontsize=8.5)
    for text in legend.get_texts():
        text.set_color('#E2E8F0')

    plt.tight_layout()
    plt.savefig(path, dpi=200, facecolor=fig.get_facecolor(), edgecolor='none')
    plt.close()
    return path

def agregar_slides_optimizacion(prs):
    """Inserta las diapositivas de Optimización SQL en la presentación."""
    img_optimizacion = generar_grafico_optimizacion()

    # --------------------------------------------------------------------------
    # SLIDE 8: OPTIMIZACIÓN DE CONSULTAS SQL
    # --------------------------------------------------------------------------
    s8 = create_base_slide(prs, "Afinamiento de Rendimiento", "Análisis de Planes de Ejecución y Métricas de E/S")

    add_card(s8, Inches(0.8), Inches(1.65), Inches(5.7), Inches(5.1), "Fundamentos de EXPLAIN ANALYZE")
    tb = s8.shapes.add_textbox(Inches(1.05), Inches(2.25), Inches(5.2), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True
    add_structured_item(tf, "Tiempo de Ejecución:", "Medición cronométrica real del procesamiento interno en el motor de la base de datos.", 12)
    add_structured_item(tf, "Buffers (Hit vs Read):", "Hit indica lectura directa en memoria RAM; Read representa lectura física de disco (cuello de botella de E/S).", 12)
    add_structured_item(tf, "Sequential Scan (Seq Scan):", "Recorrido exhaustivo de cada página de la tabla. Ineficiente para búsquedas selectivas sobre grandes volúmenes.", 12)
    add_structured_item(tf, "Index Scan:", "Acceso selectivo mediante estructuras balanceadas B-Tree con complejidad logarítmica O(log N).", 12)

    add_card(s8, Inches(6.8), Inches(1.65), Inches(5.7), Inches(5.1), "Técnicas de Indexación Aplicadas")
    tb = s8.shapes.add_textbox(Inches(7.05), Inches(2.25), Inches(5.2), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True
    add_structured_item(tf, "Índice B-Tree Único:", "Búsqueda directa por documento de identidad; evita el escaneo de 100,000 filas.", 12)
    add_structured_item(tf, "Índice Compuesto Multicolumna:", "Indexación simultánea de cuenta y fecha; elimina la necesidad de ordenamiento en memoria (Sort).", 12)
    add_structured_item(tf, "Índice Basado en Expresiones:", "Creación sobre funciones de transformación (ej. LOWER); neutraliza anti-patrones en el WHERE.", 12)
    add_structured_item(tf, "Auditoría con pg_stat_statements:", "Inspección acumulada para identificar las 10 consultas más lentas del sistema global.", 12)

    # --------------------------------------------------------------------------
    # SLIDE 9: BENCHMARK DE RENDIMIENTO (EXPLAIN)
    # --------------------------------------------------------------------------
    s9 = create_base_slide(prs, "Resultados Comparativos", "Impacto de la Optimización: De Sequential Scan a Index Scan")

    add_card(s9, Inches(0.8), Inches(1.65), Inches(5.5), Inches(5.1), "Métricas de Desempeño")
    tb = s9.shapes.add_textbox(Inches(1.05), Inches(2.25), Inches(5.0), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True
    add_structured_item(tf, "Caso 1 (Búsqueda por DNI):", "La latencia se redujo de 13.71 ms a 0.12 ms; las lecturas pasaron de 1,301 bloques a solo 4 bloques (mejora >110x).", 12)
    add_structured_item(tf, "Caso 2 (Cuentas y Fechas):", "El índice compuesto evitó el Quicksort sobre 400k filas; la latencia se redujo de 391.04 ms a 0.10 ms (mejora >3,900x).", 12)
    add_structured_item(tf, "Caso 3 (Uso de LOWER):", "El índice funcional resolvió la invalidación del índice estándar, reduciendo el tiempo de 71.17 ms a 0.72 ms (mejora >95x).", 12)
    add_structured_item(tf, "Eficiencia General:", "Menor contención de memoria compartida y liberación inmediata de conexiones en el pool.", 12)

    s9.shapes.add_picture(img_optimizacion, Inches(6.6), Inches(1.8), width=Inches(5.9))

if __name__ == "__main__":
    print(">>> Generando vista previa independiente del MÓDULO OPTIMIZACIÓN...")
    preview_prs = create_empty_deck()
    agregar_slides_optimizacion(preview_prs)
    preview_path = os.path.join(SCRIPT_DIR, "preview_modulo_optimizacion.pptx")
    save_deck_safe(preview_prs, preview_path)
